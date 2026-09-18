"""Designer-controlled, one-time import of organization, salespeople and dealers."""
from __future__ import annotations

import hashlib
import json
import re
import secrets
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import openpyxl
from flask import Blueprint, current_app, g, jsonify, request, send_from_directory, session
from werkzeug.exceptions import HTTPException

import lgsale_db as db

BASE = Path(__file__).resolve().parent
STORE = BASE / "runtime" / "initial_import"
MAX_FILE = 15 * 1024 * 1024
LEVELS = frozenset(("一般店", "DC店", "專售店", "AC店", "批店", "失聯店"))
OPTIONAL = ("shortName", "level", "contactName", "mobilePhone", "companyPhone", "postalCode", "streetAddress")
DEALER_FIELDS = ("org", "code", "name", "owner", *OPTIONAL)
STAFF_FIELDS = ("org", "number", "name")
LIMITS = {"code": 30, "name": 150, "shortName": 150, "contactName": 100,
          "mobilePhone": 30, "companyPhone": 30, "postalCode": 20, "streetAddress": 500}
LABELS = {"code": "TWCode", "name": "經銷商名稱", "owner": "負責業務", "shortName": "經銷商簡稱",
          "level": "等級", "contactName": "聯絡人", "mobilePhone": "手機", "companyPhone": "公司電話",
          "postalCode": "郵遞區號", "streetAddress": "地址"}
bp = Blueprint("initial_import", __name__)


def cell(value):
    if value is None:
        return ""
    return str(value).strip()


def get_workbook(path):
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("無法讀取 Excel 檔案，請確認是有效的 .xlsx") from exc


def inspect(path):
    workbook = get_workbook(path)
    try:
        return [{"name": sheet.title, "rows": sheet.max_row,
                 "sample": [[cell(x) for x in row[:100]] for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True)]}
                for sheet in workbook]
    finally:
        workbook.close()


def organization_candidates(path, mapping):
    """Read only the Designer-selected source column before any staff/dealer mapping."""
    if not isinstance(mapping, dict):
        raise ValueError("請先選擇處所工作表與欄位")
    workbook = get_workbook(path)
    try:
        sheet = workbook[mapping.get("orgSheet")] if mapping.get("orgSheet") in workbook.sheetnames else None
        if sheet is None:
            raise ValueError("請選擇有效的處所工作表")
        header, column = mapping.get("orgHeader"), mapping.get("orgColumn")
        if type(header) is not int or header < 1 or header > min(sheet.max_row, 20):
            raise ValueError("處所標題列須在工作表前 20 列")
        if type(column) is not int or column < 0 or column >= sheet.max_column:
            raise ValueError("請選擇有效的處所來源欄位")
        organizations, seen = [], set()
        for row_number, row in enumerate(sheet.iter_rows(min_row=header + 1, values_only=True), header + 1):
            name = cell(row[column]) if column < len(row) else ""
            if not name or name.upper() == "#N/A":
                continue
            if len(name) > 100:
                raise ValueError(f"處所工作表第 {row_number} 列名稱超過 100 字")
            if name.casefold() not in seen:
                organizations.append(name)
                seen.add(name.casefold())
            if row_number - header > 15000:
                raise ValueError("處所資料列超過 15000 筆")
        return organizations
    finally:
        workbook.close()


def selected_organizations(candidates, selected):
    if not isinstance(selected, list) or not selected:
        raise ValueError("請至少勾選一個要建立的處所")
    if any(not isinstance(name, str) or name not in candidates for name in selected):
        raise ValueError("所選處所不屬於來源檔")
    if len({name.casefold() for name in selected}) != len(selected):
        raise ValueError("所選處所不可重複")
    return selected


def validate_mapping(mapping, sheets):
    if not isinstance(mapping, dict) or mapping.get("dealerSheet") not in sheets or mapping.get("staffSheet") not in sheets:
        raise ValueError("請選擇經銷商與業務工作表")
    for kind, fields in (("dealer", DEALER_FIELDS), ("staff", STAFF_FIELDS)):
        sheet = sheets[mapping[kind + "Sheet"]]
        header = mapping.get(kind + "Header")
        if type(header) is not int or header < 1 or header > min(sheet.max_row, 20):
            raise ValueError("標題列須在工作表前 20 列")
        selected = mapping.get(kind + "Columns")
        if not isinstance(selected, dict) or any(key not in fields for key in selected):
            raise ValueError("欄位對應格式無效")
        for key in fields:
            index = selected.get(key)
            if index is None and kind == "dealer" and key in OPTIONAL:
                continue
            if type(index) is not int or index < 0 or index >= sheet.max_column:
                raise ValueError(f"{kind} 的 {key} 尚未選擇有效來源欄位")


def extract(sheet, header_row, columns):
    selected = {key: index for key, index in columns.items() if index is not None}
    rows = []
    for row_number, raw in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
        item = {key: cell(raw[index]) if index < len(raw) else "" for key, index in selected.items()}
        if any(item.values()):
            rows.append({"row": row_number, **item})
        if len(rows) > 15000:
            raise ValueError("資料列超過 15000 筆")
    return rows


def source_data(path, mapping):
    workbook = get_workbook(path)
    try:
        sheets = {sheet.title: sheet for sheet in workbook}
        validate_mapping(mapping, sheets)
        dealers = extract(sheets[mapping["dealerSheet"]], mapping["dealerHeader"], mapping["dealerColumns"])
        staff = extract(sheets[mapping["staffSheet"]], mapping["staffHeader"], mapping["staffColumns"])
        orgs = []
        for item in dealers:
            org = item.get("org", "")
            if org and org not in orgs and org.upper() != "#N/A":
                orgs.append(org)
        return dealers, staff, orgs
    finally:
        workbook.close()


def review_source(dealers, staff, org, hire_date):
    """Pure source validation; all mapped fields of one TWCode must agree."""
    errors = []
    if len(org) > 100:
        errors.append("處所名稱超過 100 字")
    if not re.fullmatch(r"\d{4}-\d{2}-01", hire_date):
        raise ValueError("到職日期須為所選月份的 1 日")
    try:
        datetime.strptime(hire_date, "%Y-%m-%d")
    except ValueError as exc:
        raise ValueError("到職月份無效") from exc
    staff_rows = [row for row in staff if row.get("org") == org]
    dealer_rows = [row for row in dealers if row.get("org") == org]
    if not staff_rows or not dealer_rows:
        errors.append("所選處所沒有完整的業務或經銷商資料")
    by_number = defaultdict(list)
    by_name = defaultdict(set)
    for row in staff_rows:
        number, name = row.get("number", ""), row.get("name", "")
        if not number or not name or len(number) > 30 or len(name) > 100:
            errors.append(f"業務工作表第 {row['row']} 列缺少員工編號或姓名，或內容過長")
            continue
        by_number[number.casefold()].append(row)
        by_name[name].add(number.casefold())
    employees = []
    for number, rows in by_number.items():
        names = {r["name"] for r in rows}
        if len(names) > 1:
            errors.append(f"業務員編號 {number} 對應多個姓名：" + "、".join(sorted(names)))
        employees.append({"row": rows[0]["row"], "number": rows[0]["number"],
                          "name": rows[0]["name"], "org": org, "hireDate": hire_date})
    for name, numbers in by_name.items():
        if len(numbers) > 1:
            errors.append(f"業務員姓名 {name} 對應多個編號，請在 Excel 修正")
    by_code = defaultdict(list)
    for row in dealer_rows:
        code = row.get("code", "")
        if not code or not code.isascii() or len(code) > 30:
            errors.append(f"經銷商工作表第 {row['row']} 列 TWCode 空白或無效")
            continue
        by_code[code.casefold()].append(row)
    result_dealers = []
    for code, rows in by_code.items():
        values = {}
        for key in DEALER_FIELDS:
            if key == "org":
                continue
            variants = defaultdict(list)
            for row in rows:
                variants[row.get(key, "")].append(row["row"])
            if len(variants) > 1:
                details = "；".join(f"{repr(v)}（第 {','.join(map(str, indices[:8]))} 列）" for v, indices in variants.items())
                errors.append(f"處所 {org}：TWCode {rows[0]['code']} 的 {LABELS[key]} 衝突：{details}")
            values[key] = rows[0].get(key, "")
        if not values["name"] or any(len(values.get(key, "")) > limit for key, limit in LIMITS.items()):
            errors.append(f"TWCode {rows[0]['code']} 的名稱或欄位長度無效")
        if values.get("level") and values["level"] not in LEVELS:
            errors.append(f"TWCode {rows[0]['code']} 的等級 {values['level']} 不在六種正式等級內")
        owner = values["owner"]
        if not owner or len(by_name.get(owner, ())) != 1:
            errors.append(f"TWCode {rows[0]['code']} 的負責業務 {owner or '空白'} 無法唯一對應本處所業務")
        result_dealers.append({"row": rows[0]["row"], "sourceRows": [r["row"] for r in rows],
                               **values, "employeeNo": next(iter(by_name[owner])) if len(by_name.get(owner, ())) == 1 else ""})
    return {"employees": employees, "dealers": result_dealers, "errors": errors,
            "sourceRows": len(dealer_rows)}


def review_database(cur, source, org):
    errors = list(source["errors"])
    cur.execute("SELECT OrgUnitId,OrgUnitCode,OrgUnitName,IsActive FROM dbo.OrganizationUnit WHERE OrgUnitName=%s", (org,))
    existing_org = cur.fetchall()
    if len(existing_org) > 1:
        errors.append(f"資料庫有多筆同名處所：{org}")
    if not existing_org:
        errors.append(f"處所 {org} 尚未建立，請先完成處所選擇與建立")
    if existing_org and not existing_org[0][3]:
        errors.append(f"處所 {org} 已停用，不能作為初始化匯入目標")
    cur.execute("SELECT e.EmployeeNo,e.EmployeeName,o.OrgUnitName,e.TerminationDate FROM dbo.Employee e LEFT JOIN dbo.EmployeeOrgAssignmentHistory a ON a.EmployeeId=e.EmployeeId AND a.EndDateTime IS NULL LEFT JOIN dbo.OrganizationUnit o ON o.OrgUnitId=a.OrgUnitId")
    existing_staff = {r[0].casefold(): r for r in cur.fetchall()}
    for person in source["employees"]:
        prior = existing_staff.get(person["number"].casefold())
        if prior and (prior[1] != person["name"] or prior[2] != org):
            errors.append(f"員工 {person['number']} 已存在，但姓名或處所與檔案不同")
        if prior and prior[3] is not None:
            errors.append(f"員工 {person['number']} 已離職，不能沿用作為目前負責業務")
        person["action"] = "保留" if prior else "新增"
    cur.execute("SELECT DealerCode FROM dbo.Dealer")
    existing_dealers = {r[0].casefold() for r in cur.fetchall()}
    for dealer in source["dealers"]:
        if dealer["code"].casefold() in existing_dealers:
            errors.append(f"TWCode {dealer['code']} 已存在；本處所不可重複匯入")
        dealer["action"] = "新增"
    source["errors"] = errors
    source["organization"] = {"name": org, "action": "已建立" if existing_org else "尚未建立"}
    source["counts"] = {"employees": len(source["employees"]), "dealers": len(source["dealers"]), "sourceRows": source["sourceRows"]}
    source["proof"] = hashlib.sha256(json.dumps(source, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    return source


def upload_path(token):
    if not isinstance(token, str) or not re.fullmatch(r"[0-9a-f]{48}", token):
        raise ValueError("預覽代碼無效")
    path, meta_path = STORE / f"{token}.xlsx", STORE / f"{token}.json"
    if not path.is_file() or not meta_path.is_file():
        raise ValueError("找不到上傳檔，請重新選擇")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    if meta["owner"] != session["user"]["id"]:
        raise ValueError("上傳檔不屬於目前帳戶")
    return path, meta_path, meta


def cross_organization_conflicts(reviews):
    """Catch keys that would collide when every selected organization is imported."""
    seen_dealers, seen_staff = {}, {}
    for review in reviews:
        org = review["organization"]["name"]
        for dealer in review["dealers"]:
            code = dealer["code"].casefold()
            previous = seen_dealers.setdefault(code, org)
            if previous != org:
                review["errors"].append(f"TWCode {dealer['code']} 同時出現在處所 {previous} 與 {org}")
        for person in review["employees"]:
            number = person["number"].casefold()
            previous = seen_staff.setdefault(number, org)
            if previous != org:
                review["errors"].append(f"員工編號 {person['number']} 同時出現在處所 {previous} 與 {org}")


def get_batch_review(payload, cur):
    path, _, meta = upload_path(payload.get("token"))
    selected = meta.get("selectedOrganizations")
    if not isinstance(selected, list) or not selected:
        raise ValueError("請先選擇並建立處所")
    dealers, staff, _ = source_data(path, payload.get("mapping"))
    reviews = []
    for org in selected:
        source = review_source(dealers, staff, org, payload.get("hireDate", ""))
        reviews.append(review_database(cur, source, org))
    cross_organization_conflicts(reviews)
    errors = [f"{review['organization']['name']}：{error}"
              for review in reviews for error in review["errors"]]
    counts = {"organizations": len(reviews),
              "employees": sum(review["counts"]["employees"] for review in reviews),
              "dealers": sum(review["counts"]["dealers"] for review in reviews),
              "sourceRows": sum(review["counts"]["sourceRows"] for review in reviews)}
    result = {"organizations": reviews, "errors": errors, "counts": counts}
    result["proof"] = hashlib.sha256(json.dumps(result, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    return result


@bp.before_request
def protect():
    if session.get("user", {}).get("type") != "EMPLOYEE" or not g.access.designer:
        return jsonify(error="僅 Designer 可使用初始化匯入"), 403
    if request.method == "POST" and not secrets.compare_digest(request.headers.get("X-Initial-CSRF", ""), session.get("initial_csrf", "missing")):
        return jsonify(error="頁面驗證已失效，請重新整理"), 403


@bp.errorhandler(ValueError)
def bad_input(exc):
    return jsonify(error=str(exc)), 400


@bp.errorhandler(Exception)
def server_error(exc):
    if isinstance(exc, HTTPException):
        return jsonify(error=exc.description), exc.code
    current_app.logger.exception("Initial import failed")
    return jsonify(error="初始化匯入失敗；請檢查資料庫 Migration 並重新預覽"), 503


@bp.get("/initial-import")
def page():
    return send_from_directory(BASE, "LGSale_InitialImport.html")


@bp.get("/api/initial-import/context")
def context():
    session.setdefault("initial_csrf", secrets.token_hex(32))
    return jsonify(csrf=session["initial_csrf"], name=session["user"]["name"],
                   flowVersion="organizations-batch-v1")


@bp.post("/api/initial-import/upload")
def upload():
    request.max_content_length = MAX_FILE + 1024 * 1024
    file = request.files.get("file")
    if file is None or not file.filename.lower().endswith(".xlsx"):
        raise ValueError("請選擇 .xlsx 檔案")
    raw = file.stream.read(MAX_FILE + 1)
    if len(raw) > MAX_FILE:
        raise ValueError("檔案不可超過 15 MB")
    token = secrets.token_hex(24)
    STORE.mkdir(parents=True, exist_ok=True)
    path = STORE / f"{token}.xlsx"
    path.write_bytes(raw)
    try:
        sheets = inspect(path)
    except Exception:
        path.unlink(missing_ok=True)
        raise
    (STORE / f"{token}.json").write_text(json.dumps({"owner": session["user"]["id"],
        "name": Path(file.filename).name, "hash": hashlib.sha256(raw).hexdigest()}, ensure_ascii=False), encoding="utf-8")
    return jsonify(token=token, sheets=sheets)


@bp.post("/api/initial-import/organizations")
def organizations():
    payload = request.get_json() or {}
    path, _, _ = upload_path(payload.get("token"))
    return jsonify(organizations=organization_candidates(path, payload.get("mapping")))


@bp.post("/api/initial-import/organizations/commit")
def commit_organizations():
    payload = request.get_json() or {}
    path, meta_path, meta = upload_path(payload.get("token"))
    mapping = payload.get("mapping")
    selected = selected_organizations(organization_candidates(path, mapping), payload.get("selected"))
    if payload.get("confirmed") is not True:
        raise ValueError("請先確認所選處所與建立順序")
    conn = db.connect()
    try:
        cur = conn.cursor()
        cur.execute("SET XACT_ABORT ON; SET TRANSACTION ISOLATION LEVEL SERIALIZABLE")
        cur.execute("DECLARE @r int; EXEC @r=sys.sp_getapplock @Resource='LGSale.InitialImport',@LockMode='Exclusive',@LockOwner='Transaction',@LockTimeout=10000; SELECT @r")
        while cur.description is None:
            if not cur.nextset():
                raise ValueError("無法取得匯入鎖定")
        if cur.fetchone()[0] < 0:
            raise ValueError("另一筆初始化匯入正在執行")
        result = []
        for name in selected:
            cur.execute("SELECT OrgUnitId,IsActive FROM dbo.OrganizationUnit WHERE OrgUnitName=%s", (name,))
            existing = cur.fetchall()
            if len(existing) > 1:
                raise ValueError(f"資料庫有多筆同名處所：{name}")
            if existing and not existing[0][1]:
                raise ValueError(f"處所 {name} 已停用，不能作為初始化匯入目標")
            if not existing:
                code = "ORG-" + hashlib.sha256(name.encode()).hexdigest()[:16].upper()
                cur.execute("INSERT dbo.OrganizationUnit(OrgUnitCode,OrgUnitName,IsActive) VALUES(%s,%s,1)", (code, name))
            result.append({"name": name, "action": "已存在" if existing else "新增"})
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    meta["selectedOrganizations"] = selected
    meta["organizationMapping"] = mapping
    temporary = meta_path.with_suffix(".tmp")
    temporary.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    temporary.replace(meta_path)
    return jsonify(organizations=result)


@bp.post("/api/initial-import/preview")
def preview():
    with db.connect() as conn:
        result = get_batch_review(request.get_json() or {}, conn.cursor())
    return jsonify(**result)


@bp.post("/api/initial-import/commit")
def commit():
    payload = request.get_json() or {}
    if payload.get("confirmed") is not True:
        raise ValueError("請先確認完整預覽")
    conn = db.connect()
    try:
        cur = conn.cursor()
        cur.execute("SET XACT_ABORT ON; SET TRANSACTION ISOLATION LEVEL SERIALIZABLE")
        cur.execute("DECLARE @r int; EXEC @r=sys.sp_getapplock @Resource='LGSale.InitialImport',@LockMode='Exclusive',@LockOwner='Transaction',@LockTimeout=10000; SELECT @r")
        while cur.description is None:
            if not cur.nextset():
                raise ValueError("無法取得匯入鎖定")
        if cur.fetchone()[0] < 0:
            raise ValueError("另一筆初始化匯入正在執行")
        result = get_batch_review(payload, cur)
        if result["proof"] != payload.get("proof"):
            raise ValueError("來源或資料庫已變動，請重新預覽")
        if result["errors"]:
            raise ValueError("尚有衝突，請修改 Excel 後重新上傳")
        actor = int(session["user"]["employeeId"])
        now = datetime.now().replace(microsecond=0)
        for organization in result["organizations"]:
            org = organization["organization"]["name"]
            cur.execute("SELECT OrgUnitId FROM dbo.OrganizationUnit WHERE OrgUnitName=%s AND IsActive=1", (org,))
            org_row = cur.fetchone()
            if not org_row:
                raise ValueError(f"處所 {org} 已不存在，請重新預覽")
            org_id = org_row[0]
            for person in organization["employees"]:
                if person["action"] != "新增":
                    continue
                cur.execute("INSERT dbo.Employee(EmployeeNo,EmployeeName,HireDate) OUTPUT inserted.EmployeeId VALUES(%s,%s,%s)", (person["number"], person["name"], person["hireDate"]))
                employee_id = cur.fetchone()[0]
                cur.execute("INSERT dbo.EmployeePositionHistory(EmployeeId,PositionLevel,StartDateTime,ChangeReason,CreatedByEmployeeId) VALUES(%s,'SALES',%s,N'初始化匯入',%s)", (employee_id, now, actor))
                cur.execute("INSERT dbo.EmployeeOrgAssignmentHistory(EmployeeId,OrgUnitId,StartDateTime,ChangeReason,CreatedByEmployeeId) VALUES(%s,%s,%s,N'初始化匯入',%s)", (employee_id, org_id, now, actor))
                cur.execute("INSERT dbo.UserAccount(AccountType,EmployeeId,IsLoginEnabled,AccountStatus) VALUES('EMPLOYEE',%s,1,'ACTIVE')", (employee_id,))
            cur.execute("SELECT EmployeeNo,EmployeeId FROM dbo.Employee")
            employee_ids = {r[0].casefold(): r[1] for r in cur.fetchall()}
            for dealer in organization["dealers"]:
                cur.execute("""INSERT dbo.Dealer(DealerCode,DealerName,ShortName,ContactName,MobilePhone,CompanyPhone,PostalCode,StreetAddress)
                    OUTPUT inserted.DealerId VALUES(%s,%s,%s,%s,%s,%s,%s,%s)""",
                    (dealer["code"], dealer["name"], *(dealer.get(key) or None for key in ("shortName", "contactName", "mobilePhone", "companyPhone", "postalCode", "streetAddress"))))
                dealer_id = cur.fetchone()[0]
                cur.execute("INSERT dbo.DealerAssignmentHistory(DealerId,EmployeeId,StartDateTime,ChangeReason,CreatedByEmployeeId) VALUES(%s,%s,%s,N'初始化匯入',%s)",
                            (dealer_id, employee_ids[dealer["employeeNo"]], now, actor))
                if dealer.get("level"):
                    cur.execute("INSERT dbo.DealerLevelHistory(DealerId,DealerStatus,StartDateTime,ChangeReason) VALUES(%s,%s,%s,N'初始化匯入')",
                                (dealer_id, dealer["level"], now))
        conn.commit()
        return jsonify(organizations=[item["organization"]["name"] for item in result["organizations"]],
                       counts=result["counts"])
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
