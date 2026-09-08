"""Opening inventory upload, review and transactional import. No writes during review."""
from __future__ import annotations

import hashlib
import json
import re
import secrets
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from zipfile import ZipFile, BadZipFile

import openpyxl
from flask import Blueprint, jsonify, request, session, send_from_directory, current_app
from werkzeug.exceptions import HTTPException
import lgsale_db as db

BASE = Path(__file__).resolve().parent
STORE = BASE / 'uploads' / 'opening_inventory'
bp = Blueprint('opening', __name__)
HEADERS = ['TW CODE', '簡稱', '產品別', '產品別2', '型號', '陳列', '期末', '不計', '業務員']
MAX_FILE = 15 * 1024 * 1024


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()


def month_date(value):
    if not isinstance(value, str) or not re.fullmatch(r'\d{4}-\d{2}', value):
        raise ValueError('請選擇有效的期初月份')
    return datetime.strptime(value + '-01', '%Y-%m-%d')


def acquire_import_lock(cur):
    # SQL Server EXEC may emit an update-count result before the SELECT.
    # pytds requires advancing to the row-producing result explicitly.
    cur.execute("DECLARE @result int; EXEC @result=sys.sp_getapplock @Resource='LGSale.OpeningImport',@LockMode='Exclusive',@LockOwner='Transaction',@LockTimeout=10000; SELECT @result AS LockResult;")
    while cur.description is None:
        if not cur.nextset():
            raise ValueError('無法取得匯入鎖定結果，請稍後重試')
    row=cur.fetchone()
    if row is None or row[0]<0:
        raise ValueError('另一筆匯入正在執行，請稍後再試')


def parse_workbook(path):
    """Strict source validation; totals are never imported as detail."""
    try:
        with ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Excel 解壓後過大，請使用 100 MB 以下的工作簿')
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError('無法讀取檔案，請上傳有效的 .xlsx') from exc
    try:
        if '期末' not in book.sheetnames:
            raise ValueError('找不到「期末」工作表')
        sheet = book['期末']
        if sheet.max_row > 50000 or sheet.max_column > 100:
            raise ValueError('工作表範圍過大；最多支援 50,000 列、100 欄')
        rows, errors, seen = [], [], set()
        header = False
        footer = False
        for number,cells in enumerate(sheet.iter_rows(max_col=9),1):
            values = [c.value.strip() if isinstance(c.value, str) else c.value for c in cells]
            if not header:
                if values == HEADERS:
                    header = True
                elif number > 20:
                    raise ValueError('前 20 列找不到預期的九個欄位標題')
                continue
            if not any(v is not None and v != '' for v in values):
                continue
            # Known source control/footer rows, not arbitrary malformed records.
            if values[0] == '總計':
                footer=True
                continue
            if footer and not any(values[:4]) and values[4] in (None, '', '不計'):
                continue
            if any(c.data_type in ('f', 'e') for c in cells):
                errors.append(f'第 {number} 列含公式或 Excel 錯誤，請先轉為確認過的值');continue
            a,b,c,d,e,f,g,h,i = values
            if not all(isinstance(v,str) and v for v in (a,b,c,d,e,i)):
                errors.append(f'第 {number} 列的代碼、名稱、分類、型號或業務姓名缺漏');continue
            if len(a)>30 or not a.isascii() or len(b)>150 or len(e)>50 or not e.isascii() or len(c)>100 or len(d)>100 or len(i)>100:
                errors.append(f'第 {number} 列文字長度或代碼格式超出主檔限制');continue
            a,e=a.upper(),e.upper()
            f = 0 if f in (None,'') else f
            if any(isinstance(v,bool) or not isinstance(v,(int,float)) or v != int(v) or not -2147483648 <= v <= 2147483647 for v in (f,g)) or f<0:
                errors.append(f'第 {number} 列數量必須為整數，陳列不可為負數');continue
            if h not in (0, '不計'):
                errors.append(f'第 {number} 列「不計」僅接受 0 或「不計」');continue
            key=(a.casefold(),e.casefold())
            if key in seen:
                errors.append(f'第 {number} 列的 TW CODE＋型號重複');continue
            seen.add(key)
            rows.append(dict(row=number,code=a,dealer=b,category1=c,category2=d,product=e,display=int(f),quantity=int(g),excluded=h=='不計',employee=i))
        if not header or not rows:
            errors.append('沒有可匯入的明細')
        return rows, errors
    finally:
        book.close()


def snapshot(cur, month):
    queries = {
        'dealers':'SELECT DealerId,DealerCode,DealerName FROM dbo.Dealer ORDER BY DealerId',
        'employees':'SELECT EmployeeId,EmployeeNo,EmployeeName,HireDate,TerminationDate FROM dbo.Employee ORDER BY EmployeeId',
        'products':'SELECT ProductId,ProductCode,ProductName,CategoryLevel1,CategoryLevel2 FROM dbo.Product ORDER BY ProductId',
        'assignments':'SELECT DealerAssignmentId,DealerId,EmployeeId,StartDateTime,EndDateTime FROM dbo.DealerAssignmentHistory ORDER BY DealerAssignmentId',
        'orgs':'SELECT OrgUnitId,OrgUnitName FROM dbo.OrganizationUnit WHERE IsActive=1 ORDER BY OrgUnitId',
    }
    result={}
    for key,sql in queries.items():
        cur.execute(sql);result[key]=[list(r) for r in cur.fetchall()]
    cur.execute("""SELECT d.DealerCode,p.ProductCode,i.OpeningInventoryDetailId,i.OpeningQuantity,
        COALESCE(cb.ImportBatchId,b.ImportBatchId),COALESCE(cb.OriginalFileName,b.OriginalFileName),COALESCE(cb.ImportedAt,b.ImportedAt)
        FROM dbo.MonthlyOpeningInventoryDetail i JOIN dbo.ImportBatch b ON b.ImportBatchId=i.ImportBatchId
        JOIN dbo.Dealer d ON d.DealerId=i.DealerId JOIN dbo.Product p ON p.ProductId=i.ProductId
        OUTER APPLY(SELECT TOP 1 c.ImportBatchId FROM dbo.OpeningInventoryCorrection c
            JOIN dbo.ImportBatch x ON x.ImportBatchId=c.ImportBatchId AND x.ImportStatus='Official'
            WHERE c.OpeningInventoryDetailId=i.OpeningInventoryDetailId ORDER BY c.CorrectionId DESC) lastChange
        LEFT JOIN dbo.ImportBatch cb ON cb.ImportBatchId=lastChange.ImportBatchId
        WHERE b.ImportType='OPENING_INVENTORY' AND b.DataMonth=%s AND b.ImportStatus='Official'""",(month.replace('-',''),))
    result['existing']=[list(r) for r in cur.fetchall()]
    cur.execute("SELECT p.ProductCode,e.EffectiveFromMonth,e.EffectiveToMonth,e.ExclusionReason FROM dbo.OpeningInventoryProductExclusion e JOIN dbo.Product p ON p.ProductId=e.ProductId WHERE e.ScopeType='ALL' ORDER BY e.ExclusionId")
    result['exclusions']=[list(r) for r in cur.fetchall()]
    cur.execute("SELECT COUNT(*) FROM sys.check_constraints WHERE parent_object_id=OBJECT_ID('dbo.MonthlyOpeningInventoryDetail') AND name='CK_MonthlyOpeningInventoryDetail_Quantity' AND is_disabled=0")
    result['schemaReady']=cur.fetchone()[0]==0
    return result


def build_review(rows, parse_errors, state, month, edits=None, retained=None, decisions=None):
    edits = edits or {}; retained=set(retained or [])
    effective=month_date(month); errors=list(parse_errors); notices=[]
    decisions=decisions or {}
    source_rows=rows
    existing_by_key=defaultdict(list)
    for old in state['existing']:existing_by_key[(old[0].casefold(),old[1].casefold())].append(old)
    duplicates=[];corrections=[];new_rows=[];selected_rows=[]
    for row in source_rows:
        matches=existing_by_key.get((row['code'].casefold(),row['product'].casefold()),[])
        if not matches:
            new_rows.append(row);selected_rows.append(row);continue
        if len(matches)>1:
            errors.append(f"{row['code']}／{row['product']} 已有多筆有效庫存，請先由管理員處理");continue
        old=matches[0];choice=decisions.get(str(row['row']),'')
        item=dict(row=row['row'],code=row['code'],dealer=row['dealer'],product=row['product'],detailId=old[2],oldQuantity=old[3],newQuantity=row['quantity'],batchId=old[4],fileName=old[5],importedAt=str(old[6]),decision=choice)
        duplicates.append(item)
        if choice not in ('keep','replace'):
            errors.append(f"第 {row['row']} 列庫存重複，請選擇保留原資料或以本次數量更正")
        elif choice=='replace':
            corrections.append(item);selected_rows.append(row)
    rows=selected_rows
    dm={r[1].casefold():r for r in state['dealers']};pm={r[1].casefold():r for r in state['products']}
    em=defaultdict(list)
    for e in state['employees']:em[e[2]].append(e)
    employees_by_id={r[0]:r for r in state['employees']}
    grouped=defaultdict(list)
    for row in rows:grouped[row['code']].append(row)
    dealers=[];employees=[];products=[];assignments=[];conflicts=[]
    new_em={};employee_ids={};new_dm={};new_pm={}
    employee_numbers={r[1].casefold() for r in state['employees']}
    input_numbers=set()
    for name in sorted({r['employee'] for r in rows}):
        matches=em[name]
        if len(matches)>1:
            errors.append(f'業務「{name}」有同名主檔，請先至員工主檔改名區分');continue
        if matches:
            employee_ids[name]=matches[0][0]
            if matches[0][4] is not None and matches[0][4] <= effective.date():
                errors.append(f'業務「{name}」在期初日已離職，請先確認主檔姓名')
            if matches[0][3]>effective.date():errors.append(f'業務「{name}」到職日晚於期初日，無法建立當期配對')
            continue
        edit=edits.get(name,{})
        no=str(edit.get('number','')).strip();hire=str(edit.get('hireDate','')).strip();org=str(edit.get('orgId',''));position=edit.get('position','SALES')
        valid=True
        if not no or len(no)>30:
            errors.append(f'新業務「{name}」請填寫員工編號');valid=False
        elif no.casefold() in employee_numbers or no.casefold() in input_numbers:
            errors.append(f'新業務「{name}」的員工編號重複');valid=False
        input_numbers.add(no.casefold())
        try:
            hire_date=date.fromisoformat(hire)
            if hire_date>effective.date():raise ValueError()
        except ValueError:
            errors.append(f'新業務「{name}」請填寫不晚於期初日的到職日');valid=False
        if org not in {str(o[0]) for o in state['orgs']}:
            errors.append(f'新業務「{name}」請選擇所屬處所');valid=False
        if position not in ('SALES','DIRECTOR','MANAGER'):
            errors.append(f'新業務「{name}」職級無效');valid=False
        item=dict(name=name,number=no,hireDate=hire,orgId=org,position=position,status='新增' if valid else '需修正',dealerCodes=sorted({r['code'] for r in rows if r['employee']==name}))
        employees.append(item);new_em[name]=item
    for code,items in grouped.items():
        first=items[0];match=dm.get(code.casefold())
        owners={i['employee'] for i in items}
        if len(owners)!=1:
            errors.append(f'{code} 同一經銷商有不同業務姓名');continue
        name=first['employee'];eid=employee_ids.get(name)
        if not match:
            item=dict(code=code,name=first['dealer'],employee=name,status='新增',dealerCodes=[code]);dealers.append(item);new_dm[code]=item
        histories=[h for h in state['assignments'] if match and h[1]==match[0]]
        active=[h for h in histories if h[3]<=effective and (h[4] is None or h[4]>effective)]
        if len(active)>1:
            errors.append(f'{code} 期初日配對期間重疊');continue
        if active:
            old=employees_by_id.get(active[0][2]);oldname=old[2] if old else str(active[0][2])
            if active[0][2]!=eid:
                conflicts.append(dict(code=code,name=first['dealer'],employee=name,current=oldname,retained=code in retained))
                if code not in retained:errors.append(f'{code} 檔案業務 {name} 與當期 {oldname} 不同，請確認保留原配對')
            continue
        if any(h[3]>effective for h in histories):
            errors.append(f'{code} 已有期初日之後的配對，請先處理歷史期間');continue
        assignments.append(dict(code=code,name=first['dealer'],employee=name,current='未指派',start=month+'-01 00:00:00',status='新增',dealerCodes=[code]))
    source_products=defaultdict(list)
    for row in rows:source_products[row['product']].append(row)
    for code,items in source_products.items():
        first=items[0]
        if len({(i['category1'],i['category2']) for i in items})>1:
            errors.append(f'型號 {code} 有不同分類，請先修正来源檔案');continue
        if code.casefold() in pm:
            old=pm[code.casefold()]
            if (old[3] or '',old[4] or '')!=(first['category1'],first['category2']):notices.append(f'{code} 分類與主檔不同，本次沿用主檔')
            continue
        item=dict(code=code,name=code,category1=first['category1'],category2=first['category2'],status='新增',dealerCodes=sorted({i['code'] for i in items}))
        products.append(item);new_pm[code]=item
    exclusions=[]
    for code in sorted({r['product'] for r in rows if r['excluded']}):
        existing=[e for e in state.get('exclusions',[]) if e[0].casefold()==code.casefold() and e[1]==month.replace('-','')]
        if existing and any(e[2] is not None or e[3]!='PSIRemove' for e in existing):
            errors.append(f'{code} 同月份已有不同的全通路排除設定，請先確認排除名單')
        exclusions.append(dict(code=code,scope='ALL',reason='PSIRemove',fromMonth=month.replace('-',''),toMonth=None,status='沿用' if existing else '新增'))
    if not state['schemaReady']:errors.append('資料庫尚未允許負庫存，請先由管理員完成限制調整')
    summary=dict(rows=len(source_rows),quantity=sum(r['quantity'] for r in source_rows),display=sum(r['display'] for r in source_rows),excluded=sum(r['excluded'] for r in source_rows))
    proof=digest(dict(state=state,month=month,rows=source_rows,edits=edits,retained=sorted(retained),decisions=decisions))
    return dict(rows=source_rows,importRows=new_rows,corrections=corrections,duplicates=duplicates,
        changes=dict(inserted=len(new_rows),corrected=len(corrections),skipped=sum(x['decision']=='keep' for x in duplicates)),
        dealers=dealers,employees=employees,products=products,assignments=assignments,exclusions=exclusions,conflicts=conflicts,errors=errors,notices=notices,summary=summary,month=month,proof=proof,orgs=[dict(id=o[0],name=o[1]) for o in state['orgs']],dealerOptions=[dict(code=c,name=n) for c,n in sorted({(x['code'],x['dealer']) for x in source_rows})])


def load_upload(token):
    if not isinstance(token,str) or not re.fullmatch(r'[a-f0-9]{48}',token):raise ValueError('無效的預覽識別碼')
    meta_path=STORE/(token+'.json')
    if not meta_path.exists():raise ValueError('找不到預覽，請重新選擇檔案')
    meta=json.loads(meta_path.read_text(encoding='utf-8'))
    if meta['owner']!=session['user']['id']:raise ValueError('此預覽不屬於目前帳號')
    if meta.get('cancelled'):raise ValueError('此預覽已取消，請重新選擇檔案')
    if datetime.fromisoformat(meta['created']) < datetime.now()-timedelta(hours=24):raise ValueError('預覽已超過 24 小時，請重新上傳')
    path=STORE/(token+'.xlsx')
    if hashlib.sha256(path.read_bytes()).hexdigest()!=meta['hash']:raise ValueError('來源檔案已變動，請重新上傳')
    return meta,path


@bp.before_request
def protect():
    user=session.get('user',{})
    if user.get('type')!='EMPLOYEE' or not user.get('employeeId'):
        return jsonify(error='僅員工帳號可使用期初匯入'),403
    if request.method!='GET' and not secrets.compare_digest(request.headers.get('X-Opening-CSRF',''),session.get('opening_csrf','missing')):
        return jsonify(error='頁面驗證已失效，請重新整理'),403


@bp.errorhandler(ValueError)
def bad_input(exc):return jsonify(error=str(exc)),400


@bp.errorhandler(Exception)
def server_error(exc):
    if isinstance(exc,HTTPException):return jsonify(error=exc.description),exc.code
    current_app.logger.exception('Opening inventory operation failed')
    return jsonify(error='資料庫或檔案處理失敗，尚未確認匯入完成。請重新比對或重試同一批次；系統會防止重複寫入。'),503


@bp.get('/opening-import')
def page():return send_from_directory(BASE,'LGSale_OpeningImport.html')


@bp.get('/api/opening-import/context')
def context():
    session.setdefault('opening_csrf',secrets.token_hex(32))
    response=jsonify(csrf=session['opening_csrf'],name=session['user']['name']);response.headers['Cache-Control']='no-store';return response


@bp.post('/api/opening-import/upload')
def upload():
    request.max_content_length=MAX_FILE+1024*1024
    file=request.files.get('file')
    if file is None or not file.filename.lower().endswith('.xlsx'):raise ValueError('請選擇 .xlsx 檔案')
    raw=file.stream.read(MAX_FILE+1)
    if len(raw)>MAX_FILE:raise ValueError('檔案不可超過 15 MB')
    token=secrets.token_hex(24);STORE.mkdir(parents=True,exist_ok=True)
    path=STORE/(token+'.xlsx');path.write_bytes(raw)
    try:rows,errors=parse_workbook(path)
    except Exception:
        path.unlink(missing_ok=True);raise
    filename=file.filename.replace('\\','/').split('/')[-1][:260]
    meta=dict(owner=session['user']['id'],name=filename,hash=hashlib.sha256(raw).hexdigest(),size=len(raw),created=datetime.now().isoformat())
    (STORE/(token+'.json')).write_text(json.dumps(meta,ensure_ascii=False),encoding='utf-8')
    found=re.search(r'(20\d{2})(0[1-9]|1[0-2])',filename);suggested=''
    if found:
        year,month=map(int,found.groups());suggested=f'{year+(month==12):04d}-{month%12+1:02d}'
    return jsonify(token=token,name=filename,suggestedMonth=suggested,parseErrors=errors,rowCount=len(rows))


def get_review(payload,cur):
    meta,path=load_upload(payload.get('token'));month=payload.get('month');month_date(month)
    rows,errors=parse_workbook(path)
    edits=payload.get('edits',{});retained=payload.get('retained',[])
    if not isinstance(edits,dict) or any(not isinstance(v,dict) for v in edits.values()) or not isinstance(retained,list) or any(not isinstance(v,str) for v in retained):raise ValueError('預覽編輯格式無效')
    decisions=payload.get('decisions',{})
    if not isinstance(decisions,dict) or any(v not in ('','keep','replace') for v in decisions.values()):raise ValueError('重複資料處理選項無效')
    review=build_review(rows,errors,snapshot(cur,month),month,edits,retained,decisions)
    return review,meta,path


@bp.post('/api/opening-import/preview')
def preview():
    conn=db.connect()
    try:
        review,meta,_=get_review(request.get_json() or {},conn.cursor())
        return jsonify(**review,fileName=meta['name'])
    finally:conn.close()


@bp.post('/api/opening-import/commit')
def commit():
    payload=request.get_json() or {}
    if payload.get('confirmed') is not True:raise ValueError('請先確認本批資料')
    meta,path=load_upload(payload.get('token'))
    conn=db.connect()
    try:
        cur=conn.cursor()
        cur.execute('SET XACT_ABORT ON; SET TRANSACTION ISOLATION LEVEL SERIALIZABLE;')
        acquire_import_lock(cur)
        cur.execute("SELECT ImportBatchId,SuccessRowCount,DataMonth FROM dbo.ImportBatch WHERE StoredFilePath=%s AND ImportStatus='Official'",(str(path),))
        completed=cur.fetchone()
        if completed:
            conn.rollback();return jsonify(batchId=completed[0],rows=completed[1],month=completed[2],alreadyImported=True)
        review,meta,path=get_review(payload,cur)
        if review['proof']!=payload.get('proof'):raise ValueError('資料或主檔已變動，請重新比對並確認')
        if review['errors']:raise ValueError('尚有問題未處理：'+'；'.join(review['errors'][:8]))
        if not review['importRows'] and not review['corrections']:
            conn.rollback()
            return jsonify(noChanges=True,**review['summary'],**review['changes'])
        actor=int(session['user']['employeeId']);effective=month_date(review['month']);reason='期初庫存匯入建立'
        for e in review['employees']:
            cur.execute('INSERT dbo.Employee(EmployeeNo,EmployeeName,HireDate) OUTPUT inserted.EmployeeId VALUES(%s,%s,%s)',(e['number'],e['name'],e['hireDate']));eid=cur.fetchone()[0]
            cur.execute('INSERT dbo.EmployeePositionHistory(EmployeeId,PositionLevel,StartDateTime,ChangeReason,CreatedByEmployeeId) VALUES(%s,%s,%s,%s,%s)',(eid,e['position'],e['hireDate'],reason,actor))
            cur.execute('INSERT dbo.EmployeeOrgAssignmentHistory(EmployeeId,OrgUnitId,StartDateTime,ChangeReason,CreatedByEmployeeId) VALUES(%s,%s,%s,%s,%s)',(eid,int(e['orgId']),e['hireDate'],reason,actor))
        for item in review['dealers']:
            cur.execute('INSERT dbo.Dealer(DealerCode,DealerName) VALUES(%s,%s)',(item['code'],item['name']))
        for item in review['products']:
            cur.execute('INSERT dbo.Product(ProductCode,ProductName,CategoryLevel1,CategoryLevel2) VALUES(%s,%s,%s,%s)',(item['code'],item['name'],item['category1'],item['category2']))
        cur.execute('SELECT DealerCode,DealerId FROM dbo.Dealer');dm={r[0].casefold():r[1] for r in cur.fetchall()}
        cur.execute('SELECT ProductCode,ProductId FROM dbo.Product');pm={r[0].casefold():r[1] for r in cur.fetchall()}
        cur.execute('SELECT EmployeeName,EmployeeId FROM dbo.Employee');em={r[0]:r[1] for r in cur.fetchall()}
        for item in review['assignments']:
            cur.execute('INSERT dbo.DealerAssignmentHistory(DealerId,EmployeeId,StartDateTime,ChangeReason,CreatedByEmployeeId) VALUES(%s,%s,%s,%s,%s)',(dm[item['code'].casefold()],em[item['employee']],effective,reason,actor))
        cur.execute("INSERT dbo.ImportBatch(ImportType,DataMonth,OriginalFileName,StoredFilePath,FileHash,FileSize,ImportStatus,TotalRowCount,SuccessRowCount,ImportedByEmployeeId) OUTPUT inserted.ImportBatchId VALUES('OPENING_INVENTORY',%s,%s,%s,%s,%s,'Official',%s,%s,%s)",(review['month'].replace('-',''),meta['name'],str(path),meta['hash'],meta['size'],len(review['rows']),len(review['importRows'])+len(review['corrections']),actor))
        batch=cur.fetchone()[0]
        if review['importRows']:
            cur.executemany('INSERT dbo.MonthlyOpeningInventoryDetail(ImportBatchId,SourceRowNumber,DealerId,ProductId,OpeningQuantity) VALUES(%s,%s,%s,%s,%s)',[(batch,x['row'],dm[x['code'].casefold()],pm[x['product'].casefold()],x['quantity']) for x in review['importRows']])
        for item in review['corrections']:
            cur.execute('INSERT dbo.OpeningInventoryCorrection(ImportBatchId,OpeningInventoryDetailId,SourceRowNumber,PreviousQuantity,NewQuantity) VALUES(%s,%s,%s,%s,%s)',(batch,item['detailId'],item['row'],item['oldQuantity'],item['newQuantity']))
            cur.execute('UPDATE dbo.MonthlyOpeningInventoryDetail SET OpeningQuantity=%s WHERE OpeningInventoryDetailId=%s AND OpeningQuantity=%s',(item['newQuantity'],item['detailId'],item['oldQuantity']))
            if cur.rowcount!=1:raise ValueError('原庫存已變動，請重新比對；本批未寫入')
        for item in review['exclusions']:
            if item['status']=='新增':
                cur.execute("INSERT dbo.OpeningInventoryProductExclusion(ProductId,ScopeType,DealerId,EffectiveFromMonth,EffectiveToMonth,ExclusionReason,CreatedByEmployeeId) VALUES(%s,'ALL',NULL,%s,NULL,'PSIRemove',%s)",(pm[item['code'].casefold()],item['fromMonth'],actor))
        cur.execute('SELECT COUNT(*),SUM(CAST(OpeningQuantity AS bigint)) FROM dbo.MonthlyOpeningInventoryDetail WHERE ImportBatchId=%s',(batch,));control=cur.fetchone();s=review['summary']
        if (control[0],control[1] or 0)!=(len(review['importRows']),sum(x['quantity'] for x in review['importRows'])):raise ValueError('匯入數量核對失敗，整批回滾')
        conn.commit()
        return jsonify(batchId=batch,**s,**review['changes'],exclusionsCreated=sum(x['status']=='新增' for x in review['exclusions']),created={key:len(review[key]) for key in ('dealers','employees','products','assignments')})
    except Exception:
        conn.rollback();raise
    finally:conn.close()


@bp.post('/api/opening-import/cancel')
def cancel():
    token=(request.get_json() or {}).get('token')
    meta,_=load_upload(token)
    meta['cancelled']=True
    (STORE/(token+'.json')).write_text(json.dumps(meta,ensure_ascii=False),encoding='utf-8')
    return jsonify(cancelled=True)
