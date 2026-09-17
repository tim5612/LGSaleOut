"""SaleIn workbook upload, review and transactional SellInTransaction import."""
from __future__ import annotations

import hashlib
import json
import secrets
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile, ZipFile

import openpyxl
from flask import Blueprint, current_app, g, jsonify, request, send_from_directory, session
from werkzeug.exceptions import HTTPException

import lgsale_db as db


BASE = Path(__file__).resolve().parent
STORE = BASE / "uploads" / "sell_in"
MAX_FILE = 15 * 1024 * 1024
bp = Blueprint("sellin", __name__)


def can_remove_batch():
    access = getattr(g, "access", None)
    return bool(access and (access.role == "ADMIN" or access.designer))

REQUIRED_HEADERS = (
    "Category", "Order No", "Item", "Sold-to", "Model(Prefix)",
    "Invoiced Qty", "Order Date", "Billing Date",
)
OPTIONAL_HEADERS = ("Invoice No", "Invoice Date")


def _text(value, label, row_number, max_length):
    value = "" if value is None else str(value).strip()
    if not value:
        raise ValueError(f"第 {row_number} 列「{label}」不可空白")
    if len(value) > max_length:
        raise ValueError(f"第 {row_number} 列「{label}」超過 {max_length} 字元")
    return value


def _optional_text(value, label, row_number, max_length):
    if value is None or str(value).strip() == "":
        return None
    value = str(value).strip()
    if len(value) > max_length:
        raise ValueError(f"第 {row_number} 列「{label}」超過 {max_length} 字元")
    return value


def _date(value, label, row_number, optional=False):
    if value in (None, ""):
        if optional:
            return None
        raise ValueError(f"第 {row_number} 列「{label}」不可空白")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y%m%d"):
            try:
                return datetime.strptime(value[:10], fmt).date()
            except ValueError:
                pass
    raise ValueError(f"第 {row_number} 列「{label}」不是有效日期")


def _quantity(value, category, row_number):
    if isinstance(value, bool) or value in (None, ""):
        raise ValueError(f"第 {row_number} 列「Invoiced Qty」不可空白")
    try:
        quantity = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError(f"第 {row_number} 列「Invoiced Qty」不是有效數量") from None
    if not quantity.is_finite() or quantity == 0 or abs(quantity) >= Decimal("1000000000000000"):
        raise ValueError(f"第 {row_number} 列「Invoiced Qty」必須是非零有效數量")
    if quantity.as_tuple().exponent < -3:
        raise ValueError(f"第 {row_number} 列「Invoiced Qty」最多三位小數")
    if category == "Sales" and quantity < 0:
        raise ValueError(f"第 {row_number} 列銷貨數量不可為負數")
    if category == "Return" and quantity > 0:
        raise ValueError(f"第 {row_number} 列退貨數量必須為負數")
    return quantity


def parse_workbook(path):
    """Read only the fields used by SellInTransaction and validate source rows."""
    try:
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Excel 解壓後過大，請使用 100 MB 以下的工作簿")
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("無法讀取檔案，請上傳有效的 .xlsx") from exc
    try:
        if "Data" not in book.sheetnames:
            raise ValueError("找不到「Data」工作表")
        sheet = book["Data"]
        if sheet.max_row is None or sheet.max_column is None:
            # The worksheet dimension record is optional in valid XLSX files.
            # Recalculate it so read-only parsing also accepts generated files.
            sheet.calculate_dimension(force=True)
        if sheet.max_row > 50000 or sheet.max_column > 150:
            raise ValueError("工作表範圍過大；最多支援 50,000 列、150 欄")
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in header_cells]
        positions = {}
        for header in REQUIRED_HEADERS + OPTIONAL_HEADERS:
            matches = [index for index, value in enumerate(headers) if value == header]
            if header in REQUIRED_HEADERS and len(matches) != 1:
                raise ValueError(f"第一列必須且只能有一個「{header}」欄位")
            if len(matches) > 1:
                raise ValueError(f"第一列的「{header}」欄位重複")
            positions[header] = matches[0] if matches else None

        rows, errors, seen = [], [], set()
        required_indexes = [positions[name] for name in REQUIRED_HEADERS]
        optional_indexes = [positions[name] for name in OPTIONAL_HEADERS if positions[name] is not None]
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
            if not any(cells[index].value not in (None, "") for index in required_indexes):
                continue
            if any(cells[index].data_type in ("f", "e") for index in required_indexes + optional_indexes):
                errors.append(f"第 {row_number} 列的必要欄位含公式或 Excel 錯誤")
                continue
            try:
                get = lambda name: cells[positions[name]].value if positions[name] is not None else None
                category = _text(get("Category"), "Category", row_number, 20)
                if category not in ("Sales", "Return"):
                    raise ValueError(f"第 {row_number} 列「Category」僅接受 Sales 或 Return")
                order_no = _text(get("Order No"), "Order No", row_number, 50)
                item_no = _text(get("Item"), "Item", row_number, 20)
                dealer_code = _text(get("Sold-to"), "Sold-to", row_number, 30).upper()
                product_code = _text(get("Model(Prefix)"), "Model(Prefix)", row_number, 50).upper()
                key = (order_no.casefold(), item_no.casefold())
                if key in seen:
                    raise ValueError(f"第 {row_number} 列的 Order No＋Item 在檔案內重複")
                seen.add(key)
                rows.append({
                    "row": row_number,
                    "category": category,
                    "orderNo": order_no,
                    "itemNo": item_no,
                    "dealerCode": dealer_code,
                    "productCode": product_code,
                    "quantity": _quantity(get("Invoiced Qty"), category, row_number),
                    "orderDate": _date(get("Order Date"), "Order Date", row_number),
                    "billingDate": _date(get("Billing Date"), "Billing Date", row_number),
                    "invoiceNo": _optional_text(get("Invoice No"), "Invoice No", row_number, 50),
                    "invoiceDate": _date(get("Invoice Date"), "Invoice Date", row_number, optional=True),
                })
            except ValueError as exc:
                errors.append(str(exc))
        if not rows:
            errors.append("沒有可辨識的 SaleIn 明細")
        return rows, errors
    finally:
        book.close()


def _json_value(value):
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral() else float(value)
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def build_review(rows, parse_errors, cur):
    cur.execute("SELECT DealerCode,DealerId,DealerName FROM dbo.Dealer")
    dealers = {str(row[0]).casefold(): (row[1], row[2]) for row in cur.fetchall()}
    cur.execute("SELECT ProductCode,ProductId,ProductName FROM dbo.Product")
    products = {str(row[0]).casefold(): (row[1], row[2]) for row in cur.fetchall()}
    cur.execute("SELECT MAX(DataMonth) FROM dbo.ImportBatch WHERE ImportType='OPENING_INVENTORY' AND ImportStatus='Official'")
    opening_result = cur.fetchone()
    opening_month = str(opening_result[0]) if opening_result and opening_result[0] else None
    opening_start = datetime.strptime(opening_month + "01", "%Y%m%d").date() if opening_month else None
    # Limit duplicate lookup to the orders in this workbook; the transaction
    # table is expected to grow continuously over the system's lifetime.
    existing = set()
    order_numbers = list(dict.fromkeys(row["orderNo"] for row in rows))
    for start in range(0, len(order_numbers), 500):
        chunk = order_numbers[start:start + 500]
        placeholders = ",".join(["%s"] * len(chunk))
        cur.execute(f"SELECT SalesDocumentNo,SalesDocumentItemNo FROM dbo.SellInTransaction WHERE SalesDocumentNo IN ({placeholders})", tuple(chunk))
        existing.update((str(row[0]).casefold(), str(row[1]).casefold()) for row in cur.fetchall())

    reviewed, counts = [], {"ready": 0, "newProduct": 0, "outsideDealer": 0, "duplicate": 0, "invalid": len(parse_errors), "beforeOpening": 0}
    for source in rows:
        dealer = dealers.get(source["dealerCode"].casefold())
        product = products.get(source["productCode"].casefold())
        key = (source["orderNo"].casefold(), source["itemNo"].casefold())
        if dealer is None:
            status, reason = "outsideDealer", "系統管理範圍外，跳過"
        elif key in existing:
            status, reason = "duplicate", "訂單項次已匯入"
        elif product is None:
            status, reason = "newProduct", "匯入時新增商品"
        else:
            status, reason = "ready", "可匯入"
        counts[status] += 1
        before_opening = status in ("ready", "newProduct") and opening_start is not None and source["billingDate"] < opening_start
        if before_opening:
            counts["beforeOpening"] += 1
        item = {key: _json_value(value) for key, value in source.items()}
        item.update(status=status, reason=reason,
                    beforeOpening=before_opening,
                    dealerName=dealer[1] if dealer else None,
                    productName=product[1] if product else None,
                    dealerId=dealer[0] if dealer else None,
                    productId=product[0] if product else None)
        reviewed.append(item)

    importable = [row for row in reviewed if row["status"] in ("ready", "newProduct")]
    new_products = len({row["productCode"].casefold() for row in importable if row["status"] == "newProduct"})
    sales = sum(row["category"] == "Sales" for row in importable)
    returns = len(importable) - sales
    quantity = sum(Decimal(str(row["quantity"])) for row in importable)
    data_date = max((date.fromisoformat(row["billingDate"]) for row in importable), default=None)
    proof_payload = [{key: row[key] for key in ("row", "orderNo", "itemNo", "dealerId", "productId", "quantity", "billingDate", "status", "beforeOpening")} for row in reviewed]
    proof = hashlib.sha256(json.dumps(proof_payload, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()
    return {
        "rows": reviewed,
        "parseErrors": parse_errors,
        "counts": counts,
        "summary": {"source": len(rows) + len(parse_errors), "ready": len(importable), "newProducts": new_products, "beforeOpening": counts["beforeOpening"], "openingMonth": f"{opening_month[:4]}-{opening_month[4:]}" if opening_month else None, "sales": sales, "returns": returns, "quantity": _json_value(quantity)},
        "dataDate": data_date.isoformat() if data_date else None,
        "proof": proof,
    }


def _meta_path(token):
    return STORE / f"{token}.json"


def load_upload(token):
    if not isinstance(token, str) or len(token) != 48 or any(c not in "0123456789abcdef" for c in token):
        raise ValueError("匯入預覽代碼無效，請重新選擇檔案")
    try:
        meta = json.loads(_meta_path(token).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        raise ValueError("找不到匯入預覽，請重新選擇檔案") from None
    if meta["owner"] != session["user"]["id"]:
        raise ValueError("此預覽不屬於目前帳號")
    if meta.get("cancelled"):
        raise ValueError("此預覽已取消，請重新選擇檔案")
    if datetime.fromisoformat(meta["created"]) < datetime.now() - timedelta(hours=24):
        raise ValueError("預覽已超過 24 小時，請重新上傳")
    path = STORE / f"{token}.xlsx"
    if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != meta["hash"]:
        raise ValueError("來源檔案已變動，請重新上傳")
    return meta, path


def acquire_import_lock(cur):
    cur.execute("DECLARE @result int; EXEC @result=sys.sp_getapplock @Resource='LGSale.SellInImport',@LockMode='Exclusive',@LockOwner='Transaction',@LockTimeout=10000; SELECT @result AS LockResult;")
    while cur.description is None:
        if not cur.nextset():
            raise ValueError("無法取得匯入鎖定結果，請稍後重試")
    row = cur.fetchone()
    if row is None or row[0] < 0:
        raise ValueError("另一筆 SaleIn 匯入正在執行，請稍後再試")


@bp.before_request
def protect():
    user = session.get("user", {})
    if user.get("type") != "EMPLOYEE" or not user.get("employeeId"):
        return jsonify(error="僅員工帳號可使用 SaleIn 匯入"), 403
    if request.method != "GET" and not secrets.compare_digest(request.headers.get("X-SellIn-CSRF", ""), session.get("sellin_csrf", "missing")):
        return jsonify(error="頁面驗證已失效，請重新整理"), 403


@bp.errorhandler(ValueError)
def bad_input(exc):
    return jsonify(error=str(exc)), 400


@bp.errorhandler(Exception)
def server_error(exc):
    if isinstance(exc, HTTPException):
        return jsonify(error=exc.description), exc.code
    current_app.logger.exception("SaleIn import failed")
    return jsonify(error="資料庫或檔案處理失敗，尚未確認匯入完成。請重新預覽；系統會防止重複寫入。"), 503


@bp.get("/sellin-import")
def page():
    return send_from_directory(BASE, "LGSale_SellInImport.html")


@bp.get("/api/sellin-import/context")
def context():
    session.setdefault("sellin_csrf", secrets.token_hex(32))
    conn = db.connect()
    try:
        cur = conn.cursor()
        cur.execute("""SELECT TOP (10) ImportBatchId,OriginalFileName,ImportedAt,SuccessRowCount,ErrorRowCount,DataDate,
                              TotalRowCount-SuccessRowCount-ErrorRowCount
                       FROM dbo.ImportBatch WHERE ImportType='SELL_IN' AND ImportStatus='Official' ORDER BY ImportBatchId DESC""")
        history = [{"batchId": row[0], "fileName": row[1], "importedAt": row[2].isoformat(sep=" ") if row[2] else "", "rows": row[3], "issues": row[4], "dataDate": row[5].isoformat() if row[5] else "", "removed": row[6]} for row in cur.fetchall()]
        response = jsonify(csrf=session["sellin_csrf"], name=session["user"]["name"], history=history, canRemove=can_remove_batch())
        response.headers["Cache-Control"] = "no-store"
        return response
    finally:
        conn.close()


@bp.get("/api/sellin-import/batches/<int:batch_id>/rows")
def batch_rows(batch_id):
    page = request.args.get("page", "1")
    if not page.isdecimal() or not 1 <= int(page) <= 100000:
        raise ValueError("頁碼無效")
    page = int(page)
    size = 50
    conn = db.connect()
    try:
        cur = conn.cursor()
        cur.execute("""SELECT ImportBatchId,SuccessRowCount FROM dbo.ImportBatch
                       WHERE ImportBatchId=%s AND ImportType='SELL_IN' AND ImportStatus='Official'""", (batch_id,))
        batch = cur.fetchone()
        if not batch:
            return jsonify(error="找不到此 SaleIn 匯入批次"), 404
        cur.execute("""SELECT t.SellInTransactionId,t.SourceRowNumber,t.SalesDocumentNo,t.SalesDocumentItemNo,
                              d.DealerCode,d.DealerName,p.ProductCode,p.ProductName,
                              t.BillingDate,t.Quantity,t.TransactionType
                       FROM dbo.SellInTransaction t
                       JOIN dbo.Dealer d ON d.DealerId=t.DealerId
                       JOIN dbo.Product p ON p.ProductId=t.ProductId
                       WHERE t.ImportBatchId=%s
                       ORDER BY t.SourceRowNumber,t.SellInTransactionId
                       OFFSET %s ROWS FETCH NEXT %s ROWS ONLY""", (batch_id, (page - 1) * size, size))
        rows = [dict(id=r[0], sourceRow=r[1], orderNo=r[2], itemNo=r[3], dealerCode=r[4], dealerName=r[5],
                     productCode=r[6], productName=r[7], billingDate=r[8].isoformat(),
                     quantity=str(r[9]), type=r[10]) for r in cur.fetchall()]
        response = jsonify(batchId=batch_id, total=batch[1], page=page, pageSize=size, rows=rows)
        response.headers["Cache-Control"] = "no-store"
        return response
    finally:
        conn.close()


@bp.post("/api/sellin-import/batches/<int:batch_id>/remove")
def remove_batch(batch_id):
    if not can_remove_batch():
        return jsonify(error="僅管理權限可移除 SaleIn 匯入批次"), 403
    if (request.get_json(silent=True) or {}).get("confirmed") is not True:
        raise ValueError("請先確認要移除此批 SaleIn 資料")
    conn = db.connect()
    try:
        cur = conn.cursor()
        cur.execute("SET XACT_ABORT ON; SET TRANSACTION ISOLATION LEVEL SERIALIZABLE;")
        acquire_import_lock(cur)
        cur.execute("""SELECT SuccessRowCount FROM dbo.ImportBatch WITH (UPDLOCK, HOLDLOCK)
                       WHERE ImportBatchId=%s AND ImportType='SELL_IN' AND ImportStatus='Official'""", (batch_id,))
        batch = cur.fetchone()
        if not batch:
            raise ValueError("找不到此 SaleIn 匯入批次，可能已被移除")
        cur.execute("DELETE FROM dbo.SellInTransaction WHERE ImportBatchId=%s", (batch_id,))
        removed = cur.rowcount
        if removed != batch[0]:
            raise ValueError("批次明細筆數與紀錄不符，未移除任何資料")
        cur.execute("DELETE FROM dbo.ImportBatch WHERE ImportBatchId=%s AND ImportType='SELL_IN'", (batch_id,))
        if cur.rowcount != 1:
            raise ValueError("批次移除失敗，未移除任何資料")
        conn.commit()
        return jsonify(batchId=batch_id, removed=removed)
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@bp.post("/api/sellin-import/batches/<int:batch_id>/remove-rows")
def remove_batch_rows(batch_id):
    if not can_remove_batch():
        return jsonify(error="僅管理權限可移除 SaleIn 匯入資料"), 403
    payload = request.get_json(silent=True) or {}
    ids = payload.get("ids")
    if payload.get("confirmed") is not True or not isinstance(ids, list) or not 1 <= len(ids) <= 1000:
        raise ValueError("請先確認並選取 1 至 1000 筆資料")
    if any(type(item) is not int or item <= 0 for item in ids) or len(set(ids)) != len(ids):
        raise ValueError("選取資料識別碼無效")
    conn = db.connect()
    try:
        cur = conn.cursor()
        cur.execute("SET XACT_ABORT ON; SET TRANSACTION ISOLATION LEVEL SERIALIZABLE;")
        acquire_import_lock(cur)
        cur.execute("""SELECT SuccessRowCount FROM dbo.ImportBatch WITH (UPDLOCK, HOLDLOCK)
                       WHERE ImportBatchId=%s AND ImportType='SELL_IN' AND ImportStatus='Official'""", (batch_id,))
        batch = cur.fetchone()
        if not batch or batch[0] < len(ids):
            raise ValueError("批次已變動，請重新載入明細")
        placeholders = ",".join(["%s"] * len(ids))
        cur.execute(f"""DELETE FROM dbo.SellInTransaction
                         WHERE ImportBatchId=%s AND SellInTransactionId IN ({placeholders})""", (batch_id, *ids))
        if cur.rowcount != len(ids):
            raise ValueError("部分選取資料已變動或不屬於此批次，未移除任何資料")
        cur.execute("""UPDATE dbo.ImportBatch SET SuccessRowCount=SuccessRowCount-%s
                       WHERE ImportBatchId=%s AND ImportType='SELL_IN' AND ImportStatus='Official'
                         AND SuccessRowCount>=%s""", (len(ids), batch_id, len(ids)))
        if cur.rowcount != 1:
            raise ValueError("批次筆數更新失敗，未移除任何資料")
        conn.commit()
        return jsonify(batchId=batch_id, removed=len(ids), remaining=batch[0] - len(ids))
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@bp.post("/api/sellin-import/upload")
def upload():
    request.max_content_length = MAX_FILE + 1024 * 1024
    file = request.files.get("file")
    if file is None or not file.filename.lower().endswith(".xlsx"):
        raise ValueError("請選擇 .xlsx 檔案")
    raw = file.stream.read(MAX_FILE + 1)
    if len(raw) > MAX_FILE:
        raise ValueError("檔案不可超過 15 MB")
    token = secrets.token_hex(24)
    STORE.mkdir(parents=True, exist_ok=True)
    path = STORE / f"{token}.xlsx"
    path.write_bytes(raw)
    try:
        rows, errors = parse_workbook(path)
    except Exception:
        path.unlink(missing_ok=True)
        raise
    filename = file.filename.replace("\\", "/").split("/")[-1][:260]
    meta = {"owner": session["user"]["id"], "name": filename, "hash": hashlib.sha256(raw).hexdigest(), "size": len(raw), "created": datetime.now().isoformat()}
    _meta_path(token).write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    return jsonify(token=token, name=filename, sourceRows=len(rows) + len(errors))


def get_review(token, cur):
    meta, path = load_upload(token)
    rows, errors = parse_workbook(path)
    review = build_review(rows, errors, cur)
    return review, meta, path


@bp.post("/api/sellin-import/preview")
def preview():
    token = (request.get_json() or {}).get("token")
    conn = db.connect()
    try:
        review, meta, _ = get_review(token, conn.cursor())
        return jsonify(**review, fileName=meta["name"])
    finally:
        conn.close()


@bp.post("/api/sellin-import/commit")
def commit():
    payload = request.get_json() or {}
    if payload.get("confirmed") is not True:
        raise ValueError("請先確認本批資料")
    meta, path = load_upload(payload.get("token"))
    conn = db.connect()
    try:
        cur = conn.cursor()
        cur.execute("SET XACT_ABORT ON; SET TRANSACTION ISOLATION LEVEL SERIALIZABLE;")
        acquire_import_lock(cur)
        cur.execute("""SELECT TOP (1) ImportBatchId,SuccessRowCount,ErrorRowCount FROM dbo.ImportBatch
                       WHERE ImportType='SELL_IN' AND FileHash=%s AND ImportStatus='Official'
                         AND NOT EXISTS (SELECT 1 FROM dbo.ImportBatch WHERE ImportType='SELL_IN'
                           AND FileHash=%s AND ImportStatus='Official'
                           AND TotalRowCount>SuccessRowCount+ErrorRowCount)
                       ORDER BY ImportBatchId DESC""", (meta["hash"], meta["hash"]))
        completed = cur.fetchone()
        if completed:
            conn.rollback()
            return jsonify(batchId=completed[0], rows=completed[1], issues=completed[2], alreadyImported=True)
        review, meta, path = get_review(payload.get("token"), cur)
        if review["proof"] != payload.get("proof"):
            raise ValueError("資料或主檔已變動，請重新預覽並確認")
        if review["summary"]["beforeOpening"] and payload.get("confirmedHistorical") is not True:
            raise ValueError("本批包含期初月份前的 SaleIn，請先確認歷史進貨資料仍要匯入")
        ready = [row for row in review["rows"] if row["status"] in ("ready", "newProduct")]
        if not ready:
            raise ValueError("本批沒有可匯入資料")
        new_product_codes = list(dict.fromkeys(row["productCode"] for row in ready if row["status"] == "newProduct"))
        for product_code in new_product_codes:
            cur.execute("INSERT dbo.Product(ProductCode,ProductName) OUTPUT inserted.ProductId VALUES(%s,%s)", (product_code, product_code))
            product_id = cur.fetchone()[0]
            for row in ready:
                if row["productCode"].casefold() == product_code.casefold():
                    row["productId"] = product_id
        actor = int(session["user"]["employeeId"])
        errors = review["summary"]["source"] - len(ready)
        error_summary = json.dumps({"counts": review["counts"], "parseErrors": review["parseErrors"][:100]}, ensure_ascii=False)
        cur.execute("""INSERT dbo.ImportBatch
            (ImportType,DataDate,OriginalFileName,StoredFilePath,FileHash,FileSize,ImportStatus,TotalRowCount,SuccessRowCount,ErrorRowCount,ErrorSummary,ImportedByEmployeeId)
            OUTPUT inserted.ImportBatchId
            VALUES('SELL_IN',%s,%s,%s,%s,%s,'Official',%s,%s,%s,%s,%s)""",
                    (review["dataDate"], meta["name"], str(path), meta["hash"], meta["size"], review["summary"]["source"], len(ready), errors, error_summary, actor))
        batch_id = cur.fetchone()[0]
        cur.executemany("""INSERT dbo.SellInTransaction
            (ImportBatchId,SourceRowNumber,DealerId,ProductId,SalesDocumentNo,SalesDocumentItemNo,InvoiceNo,OrderDate,BillingDate,InvoiceDate,InventoryEffectiveDate,Quantity,TransactionType,TransactionStatus,ReviewStatus)
            VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'VALID','APPROVED')""",
                        [(batch_id, row["row"], row["dealerId"], row["productId"], row["orderNo"], row["itemNo"], row["invoiceNo"], row["orderDate"], row["billingDate"], row["invoiceDate"], row["billingDate"], row["quantity"], "SALE" if row["category"] == "Sales" else "RETURN") for row in ready])
        cur.execute("SELECT COUNT(*),COALESCE(SUM(Quantity),0) FROM dbo.SellInTransaction WHERE ImportBatchId=%s", (batch_id,))
        control = cur.fetchone()
        expected_quantity = sum(Decimal(str(row["quantity"])) for row in ready)
        if control[0] != len(ready) or Decimal(str(control[1])) != expected_quantity:
            raise ValueError("匯入筆數或數量核對失敗，整批回滾")
        conn.commit()
        return jsonify(batchId=batch_id, rows=len(ready), issues=errors, productsCreated=len(new_product_codes), sales=review["summary"]["sales"], returns=review["summary"]["returns"], quantity=review["summary"]["quantity"], dataDate=review["dataDate"])
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


@bp.post("/api/sellin-import/cancel")
def cancel():
    token = (request.get_json() or {}).get("token")
    meta, _ = load_upload(token)
    meta["cancelled"] = True
    _meta_path(token).write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    return jsonify(cancelled=True)
