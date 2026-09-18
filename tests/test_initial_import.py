import openpyxl

from lgsale_initial_import import (cross_organization_conflicts, organization_candidates,
                                   selected_organizations, source_data, review_source)


def sample_workbook(tmp_path, conflict=False):
    book = openpyxl.Workbook()
    dealer = book.active
    dealer.title = "門市"
    dealer.append(["處所", "TWCode", "店名", "業務", "等級", "電話"])
    dealer.append(["台北營銷處", "TW001", "測試店", "小王", "DC店", "02-1234"])
    dealer.append(["台北營銷處", "TW001", "測試店", "小王", "DC店", "02-5678" if conflict else "02-1234"])
    staff = book.create_sheet("人員")
    staff.append(["單位", "編號", "姓名"])
    staff.append(["台北營銷處", "E001", "小王"])
    path = tmp_path / "pairing.xlsx"
    book.save(path)
    mapping = {"dealerSheet": "門市", "dealerHeader": 1,
               "dealerColumns": {"org": 0, "code": 1, "name": 2, "owner": 3, "level": 4, "companyPhone": 5},
               "staffSheet": "人員", "staffHeader": 1,
               "staffColumns": {"org": 0, "number": 1, "name": 2}}
    return path, mapping


def test_mapped_workbook_deduplicates_dealer(tmp_path):
    path, mapping = sample_workbook(tmp_path)
    dealers, staff, orgs = source_data(path, mapping)
    reviewed = review_source(dealers, staff, orgs[0], "2026-09-01")
    assert orgs == ["台北營銷處"]
    assert reviewed["errors"] == []
    assert len(reviewed["dealers"]) == 1
    assert reviewed["dealers"][0]["employeeNo"] == "e001"
    assert reviewed["dealers"][0]["level"] == "DC店"


def test_conflicting_selected_field_blocks_import(tmp_path):
    path, mapping = sample_workbook(tmp_path, conflict=True)
    dealers, staff, orgs = source_data(path, mapping)
    reviewed = review_source(dealers, staff, orgs[0], "2026-09-01")
    assert any("公司電話 衝突" in error and "第 2" in error for error in reviewed["errors"])


def test_same_twcode_in_unselected_organization_is_ignored(tmp_path):
    path, mapping = sample_workbook(tmp_path)
    book = openpyxl.load_workbook(path)
    book["門市"].append(["另一處所", "TW001", "不同店名", "另一業務", "一般店", "02-9999"])
    book.save(path)
    dealers, staff, _ = source_data(path, mapping)
    reviewed = review_source(dealers, staff, "台北營銷處", "2026-09-01")
    assert reviewed["errors"] == []
    assert reviewed["sourceRows"] == 2
    assert reviewed["dealers"][0]["sourceRows"] == [2, 3]


def test_same_twcode_in_two_selected_organizations_blocks_batch():
    reviews = [
        {"organization": {"name": "北處"}, "dealers": [{"code": "TW001"}],
         "employees": [{"number": "E001"}], "errors": []},
        {"organization": {"name": "南處"}, "dealers": [{"code": "TW001"}],
         "employees": [{"number": "E002"}], "errors": []},
    ]
    cross_organization_conflicts(reviews)
    assert not reviews[0]["errors"]
    assert "北處" in reviews[1]["errors"][0] and "南處" in reviews[1]["errors"][0]


def test_legacy_level_is_rejected_for_new_import(tmp_path):
    path, mapping = sample_workbook(tmp_path)
    book = openpyxl.load_workbook(path)
    book["門市"]["E2"] = "Z"
    book["門市"]["E3"] = "Z"
    book.save(path)
    dealers, staff, orgs = source_data(path, mapping)
    reviewed = review_source(dealers, staff, orgs[0], "2026-09-01")
    assert any("等級 Z 不在六種正式等級內" in error for error in reviewed["errors"])


def test_organizations_are_chosen_before_staff_and_dealers(tmp_path):
    path, _ = sample_workbook(tmp_path)
    book = openpyxl.load_workbook(path)
    dealer = book["門市"]
    dealer.append(["南部營業處", "TW002", "另一店", "小王", "一般店", ""])
    dealer.append(["#N/A", "", "", "", "", ""])
    dealer.append(["", "", "", "", "", ""])
    book.save(path)
    candidates = organization_candidates(path, {"orgSheet": "門市", "orgHeader": 1, "orgColumn": 0})
    assert candidates == ["台北營銷處", "南部營業處"]
    assert selected_organizations(candidates, ["南部營業處"]) == ["南部營業處"]
    try:
        selected_organizations(candidates, ["不在來源檔中的處所"])
    except ValueError:
        pass
    else:
        raise AssertionError("不得建立來源檔以外的處所")
