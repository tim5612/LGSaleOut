import unittest
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import MagicMock, patch

import openpyxl
from flask import Flask

import lgsale_psi as psi


def source():
    return dict(month="2026-09", currentMonth="2026-09", asOf="2026-09-14T12:00:00", fetchedAt="2026-09-14T12:00:00",
                dealers=[dict(id=i, code=f"D{i}", name=f"Dealer {i}", employeeId=i, employee=f"E{i}", orgId=1, org="Region") for i in (1, 2)],
                products=[dict(id=1, code="P1", name="Product 1", category="HA", subcategory="Fridge", price=None),
                          dict(id=2, code="P2", name="Product 2", category="TV", subcategory="OLED", price=None)],
                opening=[(1, 1, 10), (2, 1, 5), (1, 2, 2)], incoming=[(1, 1, Decimal("3"), Decimal("-1"))],
                outgoing=[(1, 1, 4)], displays=[(1, 1, 2, datetime(2026, 8, 30)), (2, 1, 1, datetime(2026, 9, 1)), (1, 2, 1, datetime(2026, 9, 1))], exclusions=[],
                displayPhotos=[(1,1,91,datetime(2026,9,2,10),"E1")])


class CalculationTests(unittest.TestCase):
    @patch("lgsale_psi.db.connect")
    def test_sql_periods_and_snapshot_contract(self, connect):
        cur = MagicMock()
        connect.return_value.__enter__.return_value.cursor.return_value = cur
        now = datetime(2026, 9, 14, 12)
        cur.fetchone.return_value = (now,)
        cur.fetchall.side_effect = [[], [], [], [], [], [], [], []]
        psi.load_source("2026-08")
        calls = cur.execute.call_args_list
        incoming = next(c for c in calls if "FROM dbo.SellInTransaction" in c.args[0])
        self.assertIn("t.TransactionStatus='VALID'", incoming.args[0])
        self.assertIn("t.ReviewStatus='APPROVED'", incoming.args[0])
        self.assertIn("b.ImportStatus='Official'", incoming.args[0])
        outgoing = next(c for c in calls if "SUM(CAST(p.SellOutQuantity" in c.args[0])
        self.assertIn("v.RecordStatus='ACTIVE'", outgoing.args[0])
        self.assertEqual(outgoing.args[1][0], now)  # late-entered August sales still count
        self.assertEqual(outgoing.args[1][1].isoformat(), "2026-08-01")
        self.assertEqual(outgoing.args[1][2].isoformat(), "2026-09-01")
        snapshot = next(c for c in calls if "WITH latest_visits" in c.args[0])
        self.assertIn("PARTITION BY v.DealerId", snapshot.args[0])
        self.assertIn("WHERE v.rn=1", snapshot.args[0])
        self.assertEqual(snapshot.args[1][0].date().isoformat(), "2026-08-31")

    def test_balance_returns_and_display(self):
        r = psi.build_report(source(), {})
        self.assertEqual(r["rows"][0]["cells"]["1"]["values"], [2, 10, 3, -1, 4, 8, 6])
        self.assertTrue(r["rows"][0]["cells"]["1"]["displayAt"].startswith("2026-08"))
        self.assertEqual(r["rows"][0]["cells"]["1"]["displayPhoto"]["id"],91)

    def test_no_snapshot_is_unknown_not_zero(self):
        s = source(); s["displays"] = []
        r = psi.build_report(s, {})
        self.assertEqual(r["rows"][0]["cells"]["1"]["values"], [None, 10, 3, -1, 4, 8, None])
        self.assertEqual(r["quality"]["missingDisplay"], 3)

    def test_salein_without_opening_creates_psi_with_zero_baseline(self):
        s = source(); s["opening"] = []
        r = psi.build_report(s, {})
        self.assertEqual(len(r["rows"]), 1)
        self.assertEqual(r["rows"][0]["cells"]["1"]["values"], [2, 0, 3, -1, 4, -2, -4])
        self.assertEqual([d["id"] for d in r["dealers"]], [1])

    def test_zero_opening_without_salein_does_not_create_psi_rows(self):
        s = source()
        s["opening"] = [(1, 1, 0)]
        s["incoming"] = []
        r = psi.build_report(s, {})
        self.assertEqual(r["rows"], [])
        self.assertEqual(r["dealers"], [])

    def test_explicit_zero_and_negative_stock(self):
        self.assertEqual(psi.metrics(dict(opening=0, display=0, outgoing=2)), [0, 0, 0, 0, 2, -2, -2])

    def test_global_and_dealer_exclusions(self):
        s = source(); s["exclusions"] = [(2, None), (1, 2)]
        r = psi.build_report(s, {})
        self.assertEqual(len(r["rows"]), 1)
        self.assertEqual([d["id"] for d in r["dealers"]], [1])
        self.assertEqual(r["quality"]["excludedPairs"], 2)

    def test_filters_and_no_duplicate_business_totals(self):
        r = psi.matrix(psi.build_report(source(), {}))
        self.assertEqual(len(r["columns"]), 5)  # 2 dealers + 2 sales totals + one scope total
        self.assertEqual(r["rows"][-1]["values"][-1], [4, 17, 3, -1, 4, 15, 11])
        filtered = psi.matrix(psi.build_report(source(), {"employee": "1", "category": "HA"}))
        self.assertEqual(filtered["rows"][-1]["values"][-1], [2, 10, 3, -1, 4, 8, 6])
        self.assertEqual(filtered["dealerCount"], 1)

    def test_partial_totals_remain_unknown(self):
        s = source(); s["displays"] = s["displays"][:1]
        r = psi.matrix(psi.build_report(s, {}))
        self.assertEqual(r["rows"][-1]["values"][-1], [None, 17, 3, -1, 4, 15, None])

    def test_search_and_empty(self):
        self.assertEqual(len(psi.build_report(source(), {"q": "p2"})["rows"]), 1)
        self.assertEqual(len(psi.build_report(source(), {"q": "dealer 2"})["dealers"]), 1)
        self.assertEqual(psi.matrix(psi.build_report(source(), {"q": "absent"}))["rows"], [])

    def test_all_levels_share_same_total(self):
        report = psi.build_report(source(), {})
        results = [psi.matrix(report, level) for level in ("dealer", "employee", "org", "company")]
        self.assertEqual([len(r["columns"]) for r in results], [5, 3, 2, 1])
        self.assertTrue(all(r["rows"][-1]["values"][-1] == results[0]["rows"][-1]["values"][-1] for r in results))
        with self.assertRaises(ValueError): psi.matrix(report, "bad")

    def test_decimal_precision(self):
        self.assertEqual(psi.sum_values([[0, 0, .1, 0, 0, .1, .1], [0, 0, .2, 0, 0, .2, .2]]), [0, 0, .3, 0, 0, .3, .3])

    def test_month_boundaries(self):
        now = datetime(2026, 9, 14, 10)
        self.assertEqual(psi.period(None, now)[0], "2026-09")
        self.assertEqual(psi.period("2025-12", now)[2], datetime(2026, 1, 1))
        self.assertEqual(psi.period("2024-02", now)[3].day, 29)
        for invalid in ("2026-13", "2026-9", "2026-10", "9999-12", "x"):
            with self.assertRaises(ValueError): psi.period(invalid, now)


class RouteTests(unittest.TestCase):
    def setUp(self):
        psi._SOURCE_CACHE.clear()
        self.app = Flask(__name__)
        self.app.secret_key = "unit-test-only"
        self.app.register_blueprint(psi.bp)
        self.client = self.app.test_client()

    def login(self, kind="EMPLOYEE"):
        with self.client.session_transaction() as s:
            s["user"] = {"type": kind, "employeeId": 1}

    def test_unauthenticated_and_dealer_denied(self):
        for path in ("/psi", "/api/psi", "/api/psi/export"):
            self.assertEqual(self.client.get(path).status_code, 403)
        self.login("DEALER")
        self.assertEqual(self.client.get("/api/psi").status_code, 403)

    @patch.object(psi, "load_source", side_effect=source)
    def test_employee_api(self, load):
        load.side_effect = lambda month: source()
        self.login()
        result = self.client.get("/api/psi?month=2026-09&level=employee")
        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json["dealerCount"], 2)
        self.assertEqual(result.headers["Cache-Control"], "no-store")
        self.assertEqual(self.client.get("/api/psi?employee=abc").status_code, 400)
        self.assertEqual(self.client.post("/api/psi").status_code, 405)

    @patch.object(psi, "load_source")
    def test_reload_reflects_sellout_edits_and_voids(self, load):
        self.login()
        s = source(); load.return_value = s
        def closing():
            return self.client.get("/api/psi?level=company&fresh=1").json["rows"][-1]["values"][-1][5]
        self.assertEqual(closing(), 15)
        s["outgoing"] = [(1, 1, 7)]
        self.assertEqual(closing(), 12)
        s["outgoing"] = []
        self.assertEqual(closing(), 19)
        self.assertEqual(load.call_count, 3)

    @patch.object(psi, "load_source")
    def test_export_matches_table_and_neutralizes_formulas(self, load):
        s = source(); s["products"][0]["code"] = '=HYPERLINK("bad")';load.return_value = s
        self.login()
        result = self.client.get("/api/psi/export?level=company")
        self.assertEqual(result.status_code, 200)
        book = openpyxl.load_workbook(BytesIO(result.data))
        sheet = book["PSI"]
        self.assertEqual(sheet.freeze_panes, "E5")
        self.assertEqual(sheet["C5"].data_type, "s")
        self.assertEqual([sheet.cell(sheet.max_row, c).value for c in range(5, 12)], [4, 17, 3, -1, 4, 15, 11])
        self.assertIn("計算說明", book.sheetnames)


if __name__ == "__main__":
    unittest.main()
